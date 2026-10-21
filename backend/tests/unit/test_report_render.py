import re
from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.reports.render import export_report, safe_cell


@pytest.mark.parametrize(
    "value", ['=HYPERLINK("x")', "+cmd", "-1+2", "@SUM(A1)", "  =1", "\ttext", "\ntext"]
)
def test_spreadsheet_formula_prefixes_are_escaped(value):
    assert safe_cell(value).startswith("'")


@pytest.mark.parametrize("format", ["csv", "xlsx", "pdf"])
def test_empty_reports_remain_valid(format):
    body = export_report(
        format,
        "Reporte vacío",
        datetime.now(UTC),
        "{}",
        {"total_operations": 0, "counts": {}, "percentages": {}},
        ["usuario"],
        [],
    )
    assert len(body) > 100
    if format == "xlsx":
        assert load_workbook(BytesIO(body))["Resultados"].max_row == 1


def test_pdf_escapes_markup_and_paginates_with_long_cells():
    rows = [{"usuario": "<script>& matrícula " + str(i), "curso": "BIO"} for i in range(150)]
    rows.append({"usuario": "x" * 12000, "curso": "BIO"})
    body = export_report(
        "pdf",
        "Matrículas",
        datetime.now(UTC),
        "{}",
        {"total_operations": 151, "counts": {}, "percentages": {}},
        ["usuario", "curso"],
        rows,
    )
    assert len(re.findall(rb"/Type\s*/Page\b", body)) > 1


@pytest.mark.parametrize("value", ["x" * 32768, "control\x01"])
def test_xlsx_rejects_unrepresentable_cells_instead_of_silent_truncation(value):
    from app.reports.render import ReportRenderError

    with pytest.raises(ReportRenderError):
        export_report(
            "xlsx",
            "Reporte",
            datetime.now(UTC),
            "{}",
            {"total_operations": 1, "counts": {}, "percentages": {}},
            ["usuario"],
            [{"usuario": value}],
        )


def test_spreadsheet_numeric_metrics_remain_numeric():
    assert safe_cell(123) == 123
    assert safe_cell(0) == 0
