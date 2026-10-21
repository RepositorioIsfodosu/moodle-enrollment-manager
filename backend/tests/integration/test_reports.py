import csv
import os
from datetime import UTC, datetime, timedelta
from io import BytesIO, StringIO
from uuid import UUID

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models import (
    AuditLog,
    CsvOperation,
    EnrollmentAttempt,
    EnrollmentJob,
    EnrollmentRecord,
    ValidationRecord,
    ValidationRun,
)
from tests.helpers import create, login
from tests.integration.test_imports import environment as environment  # noqa: PLC0414
from tests.integration.test_imports import imported

pytestmark = pytest.mark.skipif(os.getenv("RUN_INTEGRATION") != "1", reason="Dedicated services")


@pytest.fixture
def report_data(environment):
    client, headers, connection, engine, admin = environment
    file, entry = imported(client, headers)
    with Session(engine) as db:
        operation = db.scalar(
            select(CsvOperation).where(CsvOperation.import_id == UUID(entry["id"]))
        )
        operation.payload = operation.payload | {
            "metadata": {"campus": "Santo Domingo", "academic_period": "2026-2"}
        }
        run = ValidationRun(
            owner_id=admin,
            connection_id=UUID(connection),
            configuration_version=1,
            import_ids=[entry["id"]],
            state="completed",
        )
        db.add(run)
        db.flush()
        validated = ValidationRecord(
            run_id=run.id, operation_id=operation.id, sequence=1, status="valid"
        )
        db.add(validated)
        db.flush()
        job = EnrollmentJob(
            name="Reporte de prueba", owner_id=admin, connection_id=UUID(connection)
        )
        db.add(job)
        db.flush()
        for i, state in enumerate(["confirmed", "existing", "failed", "uncertain", "pending"]):
            r = EnrollmentRecord(
                job_id=job.id,
                file_id=UUID(file["id"]),
                source_row=i + 2,
                operation_key=str(i) * 64,
                user_identifier="=1+1" if i == 0 else "alice",
                course_identifier="BIO",
                role_id=5,
                state=state,
                validation_record_id=validated.id if i == 0 else None,
                error_code="timeout" if state == "uncertain" else None,
            )
            db.add(r)
            db.flush()
            if state == "confirmed":
                db.add(
                    EnrollmentAttempt(
                        record_id=r.id,
                        number=1,
                        outcome="confirmed_by_observation",
                        safe_context={"observed_at": "2026-01-02T12:34:56+00:00"},
                    )
                )
        db.commit()
        job_id = str(job.id)
    return client, headers, connection, engine, job_id, file["id"]


def test_report_counts_filters_groups_and_pagination(report_data):
    c, h, connection, _, job, file = report_data
    response = c.get("/api/v1/reports", headers=h, params={"job_id": job, "limit": 2})
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["total"] == 5 and len(data["items"]) == 2
    assert data["summary"]["counts"] == {
        "confirmed": 1,
        "existing": 1,
        "failed": 1,
        "uncertain": 1,
        "pending": 1,
    }
    assert data["summary"]["percentages"]["confirmed"] == 20
    assert c.get("/api/v1/reports", headers=h, params={"offset": 4, "limit": 2}).json()["items"]
    for kind, total in [("successful", 1), ("errors", 2), ("pending", 1)]:
        assert c.get("/api/v1/reports", headers=h, params={"kind": kind}).json()["total"] == total
    for key, value in [
        ("connection_id", connection),
        ("file_id", file),
        ("course", "BIO"),
        ("job_id", job),
    ]:
        assert c.get("/api/v1/reports", headers=h, params={key: value}).json()["total"] == 5
    for params in [
        {"campus": "Santo Domingo"},
        {"academic_period": "2026-2"},
        {"user": "=1+1"},
        {"error_code": "timeout"},
    ]:
        assert c.get("/api/v1/reports", headers=h, params=params).json()["total"] == 1
    assert (
        c.get("/api/v1/reports", headers=h, params={"kind": "course"}).json()["summary"][
            "total_operations"
        ]
        == 5
    )
    campus = c.get("/api/v1/reports", headers=h, params={"kind": "campus"}).json()
    assert any(row["recinto"] is None for row in campus["items"])
    confirmed = c.get("/api/v1/reports", headers=h, params={"kind": "successful"}).json()["items"][
        0
    ]
    assert confirmed["confirmacion_utc"] == "2026-01-02T12:34:56Z"
    future = (datetime.now(UTC) + timedelta(days=1)).isoformat()
    assert c.get("/api/v1/reports", headers=h, params={"since": future}).json()["total"] == 0


@pytest.mark.parametrize("format", ["csv", "xlsx", "pdf"])
def test_exports_are_real_files_safe_and_audited(report_data, format):
    c, h, _, engine, job, _ = report_data
    r = c.get("/api/v1/reports/export", headers=h, params={"format": format, "job_id": job})
    assert r.status_code == 200, r.text
    assert (
        "attachment;" in r.headers["content-disposition"]
        and r.headers["cache-control"] == "no-store"
    )
    if format == "csv":
        rows = list(csv.reader(StringIO(r.content.decode("utf-8-sig"))))
        assert any("'=1+1" in row for row in rows)
    elif format == "xlsx":
        book = load_workbook(BytesIO(r.content))
        sheet = book["Resultados"]
        assert sheet.max_row == 6
        assert all(cell.data_type != "f" for row in sheet for cell in row)
        assert any(cell.value == "'=1+1" for row in sheet for cell in row)
    else:
        assert r.content.startswith(b"%PDF-") and b"%%EOF" in r.content
    with Session(engine) as db:
        event = db.scalar(select(AuditLog).where(AuditLog.action == "report.exported"))
        assert event.safe_context["rows"] == 5 and event.safe_context["format"] == format


def test_export_limit_rejects_without_truncation(report_data, monkeypatch):
    c, h, *_ = report_data
    monkeypatch.setattr(get_settings(), "report_export_max_rows", 2)
    assert c.get("/api/v1/reports", headers=h).json()["total"] == 5
    for format in ["csv", "xlsx", "pdf"]:
        assert (
            c.get("/api/v1/reports/export", headers=h, params={"format": format}).status_code == 413
        )
    assert (
        c.get("/api/v1/reports/export", headers=h, params={"kind": "successful"}).status_code == 200
    )


def test_permissions_and_connection_owner_scoping(report_data):
    c, h, connection, _, _, _ = report_data
    for role in ["operador", "supervisor", "consulta"]:
        email = role + "-report@example.com"
        user = create(c, h, role, email).json()
        headers = login(c, email)
        r = c.get("/api/v1/reports", headers=headers)
        if role == "consulta":
            assert r.status_code == 403
            assert c.get("/api/v1/reports/export", headers=headers).status_code == 403
            continue
        assert r.json()["total"] == 0
        assert (
            c.post(
                f"/api/v1/moodle-connections/{connection}/grants",
                headers=h,
                json={"user_id": user["id"]},
            ).status_code
            == 201
        )
        assert c.get("/api/v1/reports", headers=headers).json()["total"] == (
            5 if role == "supervisor" else 0
        )
        assert c.get("/api/v1/reports", headers=headers, params={"kind": "audit"}).status_code == (
            200 if role == "supervisor" else 403
        )
    assert c.get("/api/v1/reports").status_code == 401


@pytest.mark.parametrize(
    "params",
    [
        {"since": "2026-01-01"},
        {"since": "2026-02-01T00:00:00Z", "until": "2026-01-01T00:00:00Z"},
        {"state": "bogus"},
        {"kind": "bogus"},
        {"kind": "audit", "user": "alice"},
        {"actor_id": "bad"},
        {"typo": "x"},
    ],
)
def test_invalid_filters_rejected(report_data, params):
    c, h, *_ = report_data
    assert c.get("/api/v1/reports", headers=h, params=params).status_code == 422


def test_audit_export_omits_context_and_record_filters(report_data):
    c, h, *_ = report_data
    r = c.get("/api/v1/reports", headers=h, params={"kind": "audit"})
    assert r.status_code == 200 and r.json()["total"] > 0
    assert r.json()["columns"] == ["id", "fecha_utc", "responsable", "accion", "entidad"]
    assert (
        c.get(
            "/api/v1/reports/export", headers=h, params={"kind": "audit", "format": "pdf"}
        ).status_code
        == 200
    )


def test_course_groups_separate_installations_and_paginate_deterministically(report_data):
    c, h, _, engine, job, file = report_data
    second = c.post(
        "/api/v1/moodle-connections",
        headers=h,
        json={
            "name": "CSV test",
            "base_url": "https://other.example.com",
            "token": "test-only-token",
        },
    )
    assert second.status_code == 201
    with Session(engine) as db:
        previous = db.get(EnrollmentJob, UUID(job))
        other = EnrollmentJob(
            name="Otro destino", owner_id=previous.owner_id, connection_id=UUID(second.json()["id"])
        )
        db.add(other)
        db.flush()
        db.add(
            EnrollmentRecord(
                job_id=other.id,
                file_id=UUID(file),
                source_row=2,
                operation_key="z" * 64,
                user_identifier="alice",
                course_identifier="BIO",
                role_id=5,
                state="confirmed",
            )
        )
        db.commit()
    whole = c.get("/api/v1/reports", headers=h, params={"kind": "course"}).json()
    assert whole["total"] == 6 and whole["summary"]["total_operations"] == 6
    assert len([r for r in whole["items"] if r["estado"] == "confirmed"]) == 2
    pages = [
        c.get(
            "/api/v1/reports", headers=h, params={"kind": "course", "offset": i, "limit": 1}
        ).json()["items"][0]
        for i in range(6)
    ]
    assert pages == whole["items"]
